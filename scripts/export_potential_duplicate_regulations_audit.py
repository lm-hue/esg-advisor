#!/usr/bin/env python3
from __future__ import annotations

import json
import re
import subprocess
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path("/Users/lucas/Documents/GitHub/esg-advisor")
OUTPUT_DIR = REPO_ROOT / "outputs"
OUTPUT_PATH = OUTPUT_DIR / f"regulation_potential_duplicates_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

SUPABASE_URL = "https://twjaqynuamghrobhdasf.supabase.co"
ANON_KEY = (
    "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9."
    "eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InR3amFxeW51YW1naHJvYmhkYXNmIiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzQ5NDc2NTcsImV4cCI6MjA5MDUyMzY1N30."
    "zLhPdZE1jo1svRpK07c0iLrMA1_ObKz6M3yYeDS65Rw"
)


def fetch_json_pages(table: str, select: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    offset = 0
    page_size = 1000

    while True:
        url = (
            f"{SUPABASE_URL}/rest/v1/{table}"
            f"?select={quote(select)}&order=id.asc&limit={page_size}&offset={offset}"
        )
        response = subprocess.check_output(
            [
                "curl",
                "-L",
                "-s",
                "--fail",
                url,
                "-H",
                f"apikey: {ANON_KEY}",
                "-H",
                f"Authorization: Bearer {ANON_KEY}",
            ],
            text=True,
        )
        page_rows = json.loads(response)
        rows.extend(page_rows)
        if len(page_rows) < page_size:
            break
        offset += page_size

    return rows


def build_regulation_identity_key(value: str | None) -> str:
    normalized = (value or "").strip()
    if not normalized:
        return ""
    normalized = normalized.lower()
    normalized = normalized.encode("ascii", "ignore").decode("ascii")
    normalized = re.sub(r"[^a-z0-9]+", " ", normalized)
    return normalized.strip()


def collect_identity_keys(record: dict[str, Any]) -> list[str]:
    return [key for key in {
        build_regulation_identity_key(record.get("title")),
        build_regulation_identity_key(record.get("formal_title")),
    } if key]


def parse_supplemental_regulations() -> list[dict[str, Any]]:
    text = (REPO_ROOT / "src/lib/regulations.ts").read_text()
    start = text.index("const SUPPLEMENTAL_REGULATIONS: RegulationRecord[] = [")
    end = text.index("function normalizeRegulationRecord(record: RegulationRecord): RegulationRecord")
    bracket_start = text.index("[", start)
    bracket_end = text.rfind("]", start, end)
    if bracket_end == -1:
        raise RuntimeError("Could not locate end of SUPPLEMENTAL_REGULATIONS array in src/lib/regulations.ts")
    section = text[bracket_start + 1 : bracket_end]

    objects: list[str] = []
    depth = 0
    current: list[str] = []
    for char in section:
        if char == "{":
            depth += 1
        if depth > 0:
            current.append(char)
        if char == "}":
            depth -= 1
            if depth == 0 and current:
                objects.append("".join(current))
                current = []

    def extract(pattern: str, value: str) -> str:
        match = re.search(pattern, value, re.S)
        return match.group(1) if match else ""

    rows: list[dict[str, Any]] = []
    for obj in objects:
        rows.append(
            {
                "id": extract(r"\bid:\s*'([^']+)'", obj),
                "title": extract(r"\btitle:\s*'([^']+)'", obj),
                "formal_title": extract(r"\bformal_title:\s*'([^']+)'", obj),
                "region": extract(r"\bregion:\s*'([^']+)'", obj),
                "source_name": extract(r"\bsource_name:\s*'([^']+)'", obj),
                "source_url": extract(r"\bsource_url:\s*(?:\n\s*)?'([^']+)'", obj),
                "official_source_url": extract(r"\bofficial_source_url:\s*(?:\n\s*)?'([^']+)'", obj),
                "policy_page_url": extract(r"\bpolicy_page_url:\s*(?:\n\s*)?'([^']+)'", obj),
            }
        )

    return [row for row in rows if row.get("id") and row.get("title")]


def describe_match(left: dict[str, Any], right: dict[str, Any]) -> str:
    labels: list[str] = []
    if build_regulation_identity_key(left.get("title")) == build_regulation_identity_key(right.get("title")) and left.get("title") and right.get("title"):
        labels.append("title=title")
    if build_regulation_identity_key(left.get("title")) == build_regulation_identity_key(right.get("formal_title")) and left.get("title") and right.get("formal_title"):
        labels.append("title=formal_title")
    if build_regulation_identity_key(left.get("formal_title")) == build_regulation_identity_key(right.get("title")) and left.get("formal_title") and right.get("title"):
        labels.append("formal_title=title")
    if build_regulation_identity_key(left.get("formal_title")) == build_regulation_identity_key(right.get("formal_title")) and left.get("formal_title") and right.get("formal_title"):
        labels.append("formal_title=formal_title")
    return ", ".join(labels)


def bool_label(value: bool) -> str:
    return "yes" if value else "no"


def style_header(row) -> None:
    fill = PatternFill("solid", fgColor="163C33")
    for cell in row:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def auto_fit(sheet) -> None:
    for idx, column_cells in enumerate(sheet.columns, start=1):
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            longest_line = max((len(line) for line in value.splitlines()), default=0)
            max_length = max(max_length, min(longest_line, 120))
        sheet.column_dimensions[get_column_letter(idx)].width = min(max(max_length + 2, 12), 60)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    regulations = fetch_json_pages(
        "regulations",
        "id,title,formal_title,region,source_name,source_url,official_source_url,policy_page_url,source_link_kind",
    )
    source_documents = fetch_json_pages(
        "regulation_source_documents",
        "regulation_id,document_type,archived_public_url",
    )
    supplemental = parse_supplemental_regulations()

    pdf_by_regulation: dict[str, bool] = defaultdict(bool)
    for document in source_documents:
        if (document.get("document_type") or "").lower() == "pdf" and (document.get("archived_public_url") or "").strip():
            pdf_by_regulation[str(document.get("regulation_id") or "")] = True

    db_by_key: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for regulation in regulations:
        for key in collect_identity_keys(regulation):
            db_by_key[key].append(regulation)

    supplemental_vs_db_rows: list[dict[str, Any]] = []
    for sup in supplemental:
        matched_ids: set[str] = set()
        for key in collect_identity_keys(sup):
            for db in db_by_key.get(key, []):
                if db["id"] in matched_ids:
                    continue
                matched_ids.add(db["id"])
                supplemental_vs_db_rows.append(
                    {
                        "Match Key": key,
                        "Match Detail": describe_match(db, sup),
                        "DB ID": db.get("id") or "",
                        "DB Title": db.get("title") or "",
                        "DB Formal Title": db.get("formal_title") or "",
                        "DB Region": db.get("region") or "",
                        "DB Source Name": db.get("source_name") or "",
                        "DB Source URL": db.get("source_url") or "",
                        "DB Official Website URL": db.get("official_source_url") or "",
                        "DB Policy Page URL": db.get("policy_page_url") or "",
                        "DB Source Link Kind": db.get("source_link_kind") or "",
                        "DB Has Archived PDF": bool_label(pdf_by_regulation[db["id"]]),
                        "Supplemental ID": sup.get("id") or "",
                        "Supplemental Title": sup.get("title") or "",
                        "Supplemental Formal Title": sup.get("formal_title") or "",
                        "Supplemental Region": sup.get("region") or "",
                        "Supplemental Source Name": sup.get("source_name") or "",
                        "Supplemental Source URL": sup.get("source_url") or "",
                        "Supplemental Official Website URL": sup.get("official_source_url") or "",
                        "Supplemental Policy Page URL": sup.get("policy_page_url") or "",
                    }
                )

    seen_pairs: set[tuple[str, str, str]] = set()
    db_internal_rows: list[dict[str, Any]] = []
    ids_by_key: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for regulation in regulations:
        for key in collect_identity_keys(regulation):
            ids_by_key[key].append(regulation)

    for key, matches in ids_by_key.items():
        unique_matches = {row["id"]: row for row in matches}.values()
        unique_matches = list(unique_matches)
        if len(unique_matches) < 2:
            continue
        for i, left in enumerate(unique_matches):
            for right in unique_matches[i + 1 :]:
                pair_key = (key, min(left["id"], right["id"]), max(left["id"], right["id"]))
                if pair_key in seen_pairs:
                    continue
                seen_pairs.add(pair_key)
                db_internal_rows.append(
                    {
                        "Match Key": key,
                        "Match Detail": describe_match(left, right),
                        "Left ID": left.get("id") or "",
                        "Left Title": left.get("title") or "",
                        "Left Formal Title": left.get("formal_title") or "",
                        "Left Region": left.get("region") or "",
                        "Left Source Name": left.get("source_name") or "",
                        "Left Source URL": left.get("source_url") or "",
                        "Left Official Website URL": left.get("official_source_url") or "",
                        "Left Policy Page URL": left.get("policy_page_url") or "",
                        "Left Source Link Kind": left.get("source_link_kind") or "",
                        "Left Has Archived PDF": bool_label(pdf_by_regulation[left["id"]]),
                        "Right ID": right.get("id") or "",
                        "Right Title": right.get("title") or "",
                        "Right Formal Title": right.get("formal_title") or "",
                        "Right Region": right.get("region") or "",
                        "Right Source Name": right.get("source_name") or "",
                        "Right Source URL": right.get("source_url") or "",
                        "Right Official Website URL": right.get("official_source_url") or "",
                        "Right Policy Page URL": right.get("policy_page_url") or "",
                        "Right Source Link Kind": right.get("source_link_kind") or "",
                        "Right Has Archived PDF": bool_label(pdf_by_regulation[right["id"]]),
                    }
                )

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary_rows = [
        ("Generated At (UTC)", datetime.now(timezone.utc).isoformat()),
        ("Live DB Regulations", len(regulations)),
        ("Supplemental Regulations Parsed", len(supplemental)),
        ("Potential DB vs Supplemental Duplicates", len(supplemental_vs_db_rows)),
        ("Potential DB Internal Duplicate Pairs", len(db_internal_rows)),
        ("Method", "Potential duplicates are detected by normalized title/formal_title identity keys, matching the frontend deduping fix."),
    ]
    for key, value in summary_rows:
        summary.append([key, value])
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 120
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    def add_sheet(name: str, rows: list[dict[str, Any]]) -> None:
        sheet = workbook.create_sheet(name)
        headers = list(rows[0].keys()) if rows else ["No rows"]
        sheet.append(headers)
        style_header(sheet[1])
        if rows:
            for row in rows:
                sheet.append([row.get(header, "") for header in headers])
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        auto_fit(sheet)

    add_sheet("DB vs Supplemental", supplemental_vs_db_rows)
    add_sheet("DB Internal", db_internal_rows)

    workbook.save(OUTPUT_PATH)
    print(json.dumps({
        "output": str(OUTPUT_PATH),
        "db_vs_supplemental_rows": len(supplemental_vs_db_rows),
        "db_internal_rows": len(db_internal_rows),
    }, indent=2))


if __name__ == "__main__":
    main()
