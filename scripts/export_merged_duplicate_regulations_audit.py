#!/usr/bin/env python3
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path("/Users/lucas/Documents/GitHub/esg-advisor")
INPUT_PATH = Path("/tmp/merge-exact-pdf-duplicates-report.json")
OUTPUT_PATH = REPO_ROOT / "outputs" / f"regulation_merged_exact_pdf_duplicates_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

SCALAR_FIELDS = [
    "formal_title",
    "summary",
    "description",
    "full_description",
    "category",
    "status",
    "effective_date",
    "source_url",
    "official_source_url",
    "policy_page_url",
    "source_link_kind",
]
ARRAY_FIELDS = ["tags", "topics"]


def normalize(value: Any) -> Any:
    if value is None:
        return ""
    return value


def unique_join(values: list[str]) -> str:
    seen: set[str] = set()
    ordered: list[str] = []
    for value in values:
        text = str(value).strip()
        if not text or text in seen:
            continue
        seen.add(text)
        ordered.append(text)
    return "\n".join(ordered)


def style_header(row) -> None:
    fill = PatternFill("solid", fgColor="163C33")
    for cell in row:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def auto_fit(sheet) -> None:
    for idx, column_cells in enumerate(sheet.columns, start=1):
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            longest_line = max((len(line) for line in value.splitlines()), default=0)
            max_length = max(max_length, min(longest_line, 120))
        sheet.column_dimensions[get_column_letter(idx)].width = min(max(max_length + 2, 12), 60)


def describe_enrichment(pair: dict[str, Any]) -> tuple[list[str], list[str], list[str]]:
    keeper = pair["keeper"]
    loser = pair["loser"]
    merged = pair["mergedRegulation"]

    scalar_changes: list[str] = []
    array_changes: list[str] = []
    all_changes: list[str] = []

    for field in SCALAR_FIELDS:
        keeper_value = normalize(keeper.get(field))
        loser_value = normalize(loser.get(field))
        merged_value = normalize(merged.get(field))

        if merged_value != keeper_value:
            if merged_value == loser_value and str(merged_value).strip():
                scalar_changes.append(field)
                all_changes.append(f"{field}: adopted loser value")
            else:
                scalar_changes.append(field)
                all_changes.append(f"{field}: changed in merged record")

    for field in ARRAY_FIELDS:
        keeper_values = [str(value) for value in (keeper.get(field) or [])]
        loser_values = [str(value) for value in (loser.get(field) or [])]
        merged_values = [str(value) for value in (merged.get(field) or [])]
        added_from_loser = [value for value in loser_values if value not in keeper_values and value in merged_values]
        if added_from_loser:
            array_changes.append(f"{field}: {', '.join(added_from_loser)}")
            all_changes.append(f"{field}: added {', '.join(added_from_loser)}")

    return scalar_changes, array_changes, all_changes


def main() -> None:
    report = json.loads(INPUT_PATH.read_text())
    pairs: list[dict[str, Any]] = report.get("merged_pairs", [])

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    details_sheet = workbook.create_sheet("Merged Pairs")
    changes_sheet = workbook.create_sheet("Field Changes")

    summary_rows = [
        ("Generated At", report.get("generated_at")),
        ("Dry Run", report.get("dry_run")),
        ("Merged Pair Count", report.get("pair_count")),
        ("Deleted Duplicate PDF Rows", sum(len(pair.get("deletedDuplicateDocuments", [])) for pair in pairs)),
        ("Deleted Storage Objects", sum(sum(1 for item in pair.get("deletedStorage", []) if item.get("deleted")) for pair in pairs)),
        ("Watchlists Remapped", sum(len(pair.get("watchlistChanges", [])) for pair in pairs)),
    ]

    summary_sheet.append(["Metric", "Value"])
    style_header(summary_sheet[1])
    for row in summary_rows:
        summary_sheet.append(row)
    auto_fit(summary_sheet)

    details_headers = [
        "Title",
        "Keeper ID",
        "Loser ID",
        "Region",
        "Reason",
        "Keeper Score",
        "Loser Score",
        "Fields Enriched From Loser",
        "Array Enrichments",
        "Deleted Duplicate PDF Document IDs",
        "Deleted Storage Paths",
        "Watchlist Remap Count",
        "Keeper Official URL",
        "Loser Official URL",
        "Merged Official URL",
        "Keeper Source URL",
        "Loser Source URL",
        "Merged Source URL",
        "Merged Tags",
        "Merged Topics",
    ]
    details_sheet.append(details_headers)
    style_header(details_sheet[1])

    changes_headers = [
        "Title",
        "Keeper ID",
        "Loser ID",
        "Change Detail",
        "Keeper Value",
        "Loser Value",
        "Merged Value",
    ]
    changes_sheet.append(changes_headers)
    style_header(changes_sheet[1])

    for pair in pairs:
        keeper = pair["keeper"]
        loser = pair["loser"]
        merged = pair["mergedRegulation"]
        scalar_changes, array_changes, all_changes = describe_enrichment(pair)

        details_sheet.append(
            [
                merged.get("title"),
                keeper.get("id"),
                loser.get("id"),
                merged.get("region"),
                pair.get("reason"),
                keeper.get("score"),
                loser.get("score"),
                ", ".join(scalar_changes) or "",
                unique_join(array_changes),
                unique_join([doc.get("id", "") for doc in pair.get("deletedDuplicateDocuments", [])]),
                unique_join([item.get("archived_storage_path", "") for item in pair.get("deletedStorage", []) if item.get("archived_storage_path")]),
                len(pair.get("watchlistChanges", [])),
                keeper.get("official_source_url"),
                loser.get("official_source_url"),
                merged.get("official_source_url"),
                keeper.get("source_url"),
                loser.get("source_url"),
                merged.get("source_url"),
                unique_join(merged.get("tags", [])),
                unique_join(merged.get("topics", [])),
            ]
        )

        for field in SCALAR_FIELDS:
            keeper_value = normalize(keeper.get(field))
            loser_value = normalize(loser.get(field))
            merged_value = normalize(merged.get(field))
            if merged_value != keeper_value:
                changes_sheet.append(
                    [
                        merged.get("title"),
                        keeper.get("id"),
                        loser.get("id"),
                        field,
                        str(keeper_value),
                        str(loser_value),
                        str(merged_value),
                    ]
                )

        for field in ARRAY_FIELDS:
            keeper_values = [str(value) for value in (keeper.get(field) or [])]
            loser_values = [str(value) for value in (loser.get(field) or [])]
            merged_values = [str(value) for value in (merged.get(field) or [])]
            added_from_loser = [value for value in loser_values if value not in keeper_values and value in merged_values]
            if added_from_loser:
                changes_sheet.append(
                    [
                        merged.get("title"),
                        keeper.get("id"),
                        loser.get("id"),
                        field,
                        ", ".join(keeper_values),
                        ", ".join(loser_values),
                        ", ".join(merged_values),
                    ]
                )

        if not all_changes:
            changes_sheet.append(
                [
                    merged.get("title"),
                    keeper.get("id"),
                    loser.get("id"),
                    "No field-level enrichment beyond keeper selection",
                    "",
                    "",
                    "",
                ]
            )

    for sheet in (summary_sheet, details_sheet, changes_sheet):
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        auto_fit(sheet)

    workbook.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
