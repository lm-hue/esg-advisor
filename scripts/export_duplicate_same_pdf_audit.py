#!/usr/bin/env python3
from __future__ import annotations

import json
import subprocess
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path("/Users/lucas/Documents/GitHub/esg-advisor")
INPUT_PATH = REPO_ROOT / "outputs" / "regulation_potential_duplicates_2026-05-12.xlsx"
OUTPUT_PATH = REPO_ROOT / "outputs" / f"regulation_duplicate_same_pdf_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

SUPABASE_URL = "https://twjaqynuamghrobhdasf.supabase.co"
ANON_KEY = (
    "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9."
    "eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InR3amFxeW51YW1naHJvYmhkYXNmIiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzQ5NDc2NTcsImV4cCI6MjA5MDUyMzY1N30."
    "zLhPdZE1jo1svRpK07c0iLrMA1_ObKz6M3yYeDS65Rw"
)


def fetch_json_pages(table: str, select: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    offset = 0
    page_size = 1000
    while True:
        url = f"{SUPABASE_URL}/rest/v1/{table}?select={quote(select)}&order=id.asc&limit={page_size}&offset={offset}"
        response = subprocess.check_output(
            [
                "curl",
                "-L",
                "-s",
                "--fail",
                url,
                "-H",
                f"apikey: {ANON_KEY}",
                "-H",
                f"Authorization: Bearer {ANON_KEY}",
            ],
            text=True,
        )
        page_rows = json.loads(response)
        rows.extend(page_rows)
        if len(page_rows) < page_size:
            break
        offset += page_size
    return rows


def load_sheet_rows(sheet_name: str) -> list[dict[str, Any]]:
    wb = load_workbook(INPUT_PATH, read_only=True)
    ws = wb[sheet_name]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append({headers[idx]: value for idx, value in enumerate(row)})
    return rows


def normalize_text(value: Any) -> str:
    return ("" if value is None else str(value)).strip()


def unique_join(values: list[str]) -> str:
    seen: set[str] = set()
    ordered: list[str] = []
    for value in values:
        normalized = normalize_text(value)
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        ordered.append(normalized)
    return "\n".join(ordered)


def style_header(row) -> None:
    fill = PatternFill("solid", fgColor="163C33")
    for cell in row:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def auto_fit(sheet) -> None:
    for idx, column_cells in enumerate(sheet.columns, start=1):
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            longest_line = max((len(line) for line in value.splitlines()), default=0)
            max_length = max(max_length, min(longest_line, 120))
        sheet.column_dimensions[get_column_letter(idx)].width = min(max(max_length + 2, 12), 60)


def collect_fingerprints(pdf_docs: list[dict[str, Any]]) -> tuple[set[str], set[str], set[str]]:
    sha_values = {normalize_text(doc.get("content_sha256")) for doc in pdf_docs if normalize_text(doc.get("content_sha256"))}
    archived_urls = {normalize_text(doc.get("archived_public_url")) for doc in pdf_docs if normalize_text(doc.get("archived_public_url"))}
    document_urls = {normalize_text(doc.get("document_url")) for doc in pdf_docs if normalize_text(doc.get("document_url"))}
    return sha_values, archived_urls, document_urls


def compare_pdf_sets(left_docs: list[dict[str, Any]], right_docs: list[dict[str, Any]]) -> tuple[str, str, list[str]]:
    left_sha, left_archived, left_document = collect_fingerprints(left_docs)
    right_sha, right_archived, right_document = collect_fingerprints(right_docs)

    shared_sha = sorted(left_sha & right_sha)
    if shared_sha:
        return "exact_same_pdf", "content_sha256", shared_sha

    shared_archived = sorted(left_archived & right_archived)
    if shared_archived:
        return "exact_same_pdf", "archived_public_url", shared_archived

    if left_docs and right_docs:
        shared_document = sorted(left_document & right_document)
        if shared_document:
            return "possible_same_pdf_source_url", "document_url", shared_document
        return "different_pdfs", "", []

    if left_docs and not right_docs:
        return "only_left_has_pdf", "", []
    if right_docs and not left_docs:
        return "only_right_has_pdf", "", []
    return "neither_side_has_pdf", "", []


def regulation_doc_summary(pdf_docs: list[dict[str, Any]]) -> tuple[str, str, str]:
    return (
        unique_join([doc.get("content_sha256") or "" for doc in pdf_docs]),
        unique_join([doc.get("archived_public_url") or "" for doc in pdf_docs]),
        unique_join([doc.get("document_url") or "" for doc in pdf_docs]),
    )


def main() -> None:
    pdf_documents = fetch_json_pages(
        "regulation_source_documents",
        "id,regulation_id,document_type,version_label,content_sha256,archived_public_url,document_url,source_url,updated_at",
    )
    pdf_by_regulation: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for document in pdf_documents:
        if normalize_text(document.get("document_type")).lower() != "pdf":
            continue
        pdf_by_regulation[normalize_text(document.get("regulation_id"))].append(document)

    db_vs_supplemental = load_sheet_rows("DB vs Supplemental")
    db_internal = load_sheet_rows("DB Internal")

    result_rows: list[dict[str, Any]] = []
    status_counter = Counter()

    for row in db_vs_supplemental:
        left_id = normalize_text(row.get("DB ID"))
        right_id = normalize_text(row.get("Supplemental ID"))
        left_docs = pdf_by_regulation.get(left_id, [])
        right_docs = pdf_by_regulation.get(right_id, [])
        status, match_basis, shared_values = compare_pdf_sets(left_docs, right_docs)
        status_counter[status] += 1
        left_hashes, left_archived, left_document = regulation_doc_summary(left_docs)
        right_hashes, right_archived, right_document = regulation_doc_summary(right_docs)
        result_rows.append(
            {
                "Duplicate Bucket": "DB vs Supplemental",
                "Match Key": row.get("Match Key") or "",
                "Match Detail": row.get("Match Detail") or "",
                "Comparison Status": status,
                "Match Basis": match_basis,
                "Shared Values": unique_join(shared_values),
                "Left ID": left_id,
                "Left Title": row.get("DB Title") or "",
                "Left Region": row.get("DB Region") or "",
                "Left Official Website URL": row.get("DB Official Website URL") or "",
                "Left Policy Page URL": row.get("DB Policy Page URL") or "",
                "Left PDF Count": len(left_docs),
                "Left PDF Hashes": left_hashes,
                "Left Archived PDF URLs": left_archived,
                "Left Document URLs": left_document,
                "Right ID": right_id,
                "Right Title": row.get("Supplemental Title") or "",
                "Right Region": row.get("Supplemental Region") or "",
                "Right Official Website URL": row.get("Supplemental Official Website URL") or "",
                "Right Policy Page URL": row.get("Supplemental Policy Page URL") or "",
                "Right PDF Count": len(right_docs),
                "Right PDF Hashes": right_hashes,
                "Right Archived PDF URLs": right_archived,
                "Right Document URLs": right_document,
            }
        )

    for row in db_internal:
        left_id = normalize_text(row.get("Left ID"))
        right_id = normalize_text(row.get("Right ID"))
        left_docs = pdf_by_regulation.get(left_id, [])
        right_docs = pdf_by_regulation.get(right_id, [])
        status, match_basis, shared_values = compare_pdf_sets(left_docs, right_docs)
        status_counter[status] += 1
        left_hashes, left_archived, left_document = regulation_doc_summary(left_docs)
        right_hashes, right_archived, right_document = regulation_doc_summary(right_docs)
        result_rows.append(
            {
                "Duplicate Bucket": "DB Internal",
                "Match Key": row.get("Match Key") or "",
                "Match Detail": row.get("Match Detail") or "",
                "Comparison Status": status,
                "Match Basis": match_basis,
                "Shared Values": unique_join(shared_values),
                "Left ID": left_id,
                "Left Title": row.get("Left Title") or "",
                "Left Region": row.get("Left Region") or "",
                "Left Official Website URL": row.get("Left Official Website URL") or "",
                "Left Policy Page URL": row.get("Left Policy Page URL") or "",
                "Left PDF Count": len(left_docs),
                "Left PDF Hashes": left_hashes,
                "Left Archived PDF URLs": left_archived,
                "Left Document URLs": left_document,
                "Right ID": right_id,
                "Right Title": row.get("Right Title") or "",
                "Right Region": row.get("Right Region") or "",
                "Right Official Website URL": row.get("Right Official Website URL") or "",
                "Right Policy Page URL": row.get("Right Policy Page URL") or "",
                "Right PDF Count": len(right_docs),
                "Right PDF Hashes": right_hashes,
                "Right Archived PDF URLs": right_archived,
                "Right Document URLs": right_document,
            }
        )

    deduped_rows: list[dict[str, Any]] = []
    seen_rows: set[tuple[str, str, str, str, str, str]] = set()
    for row in result_rows:
        row_key = (
            normalize_text(row["Duplicate Bucket"]),
            min(normalize_text(row["Left ID"]), normalize_text(row["Right ID"])),
            max(normalize_text(row["Left ID"]), normalize_text(row["Right ID"])),
            normalize_text(row["Comparison Status"]),
            normalize_text(row["Match Basis"]),
            normalize_text(row["Shared Values"]),
        )
        if row_key in seen_rows:
            continue
        seen_rows.add(row_key)
        deduped_rows.append(row)

    exact_rows = [row for row in deduped_rows if row["Comparison Status"] == "exact_same_pdf"]
    possible_rows = [row for row in deduped_rows if row["Comparison Status"] == "possible_same_pdf_source_url"]
    other_rows = [row for row in deduped_rows if row["Comparison Status"] not in {"exact_same_pdf", "possible_same_pdf_source_url"}]

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary_rows = [
        ("Generated At (UTC)", datetime.now(timezone.utc).isoformat()),
        ("Input Duplicate Workbook", str(INPUT_PATH)),
        ("Total Duplicate Pairs Checked", len(deduped_rows)),
        ("Exact Same PDF Pairs", len(exact_rows)),
        ("Possible Same PDF (same document URL)", len(possible_rows)),
        ("Different PDFs", status_counter["different_pdfs"]),
        ("Only Left Has PDF", status_counter["only_left_has_pdf"]),
        ("Only Right Has PDF", status_counter["only_right_has_pdf"]),
        ("Neither Side Has PDF", status_counter["neither_side_has_pdf"]),
    ]
    for key, value in summary_rows:
        summary.append([key, value])
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 100
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    def add_sheet(name: str, rows: list[dict[str, Any]]) -> None:
        sheet = workbook.create_sheet(name)
        headers = list(rows[0].keys()) if rows else ["No rows"]
        sheet.append(headers)
        style_header(sheet[1])
        if rows:
            for row in rows:
                sheet.append([row.get(header, "") for header in headers])
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        auto_fit(sheet)

    add_sheet("Exact Same PDF", exact_rows)
    add_sheet("Possible Same Source URL", possible_rows)
    add_sheet("Other Results", other_rows)

    workbook.save(OUTPUT_PATH)
    print(json.dumps({
        "output": str(OUTPUT_PATH),
        "total_pairs": len(deduped_rows),
        "exact_same_pdf_pairs": len(exact_rows),
        "possible_same_source_url_pairs": len(possible_rows),
        "different_pdfs": status_counter["different_pdfs"],
        "only_left_has_pdf": status_counter["only_left_has_pdf"],
        "only_right_has_pdf": status_counter["only_right_has_pdf"],
        "neither_side_has_pdf": status_counter["neither_side_has_pdf"],
    }, indent=2))


if __name__ == "__main__":
    main()
