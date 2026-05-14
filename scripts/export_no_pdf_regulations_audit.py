#!/usr/bin/env python3
from __future__ import annotations

import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
import subprocess
from typing import Any
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path("/Users/lucas/Documents/GitHub/esg-advisor")
OUTPUT_DIR = REPO_ROOT / "outputs"
OUTPUT_PATH = OUTPUT_DIR / f"regulations_without_pdf_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

SUPABASE_URL = "https://twjaqynuamghrobhdasf.supabase.co"
SERVICE_ROLE_KEY = (
    "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9."
    "eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InR3amFxeW51YW1naHJvYmhkYXNmIiwicm9sZSI6InNlcnZpY2Vfcm9sZSIs"
    "ImlhdCI6MTc3NDk0NzY1NywiZXhwIjoyMDkwNTIzNjU3fQ."
    "5fzlwaVe-6e6Q_uhWYJG0qxerbC7pGzjiSAltGQgShw"
)

REGULATION_SELECT = ",".join(
    [
        "id",
        "title",
        "formal_title",
        "region",
        "category",
        "status",
        "source_name",
        "source_url",
        "official_source_url",
        "policy_page_url",
        "source_link_kind",
        "updated_at",
    ]
)

SOURCE_DOCUMENT_SELECT = ",".join(
    [
        "id",
        "regulation_id",
        "document_type",
        "version_label",
        "source_name",
        "source_url",
        "official_source_url",
        "policy_page_url",
        "document_url",
        "archived_public_url",
        "updated_at",
    ]
)


def fetch_json_pages(table: str, select: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    offset = 0
    page_size = 1000

    while True:
        url = (
            f"{SUPABASE_URL}/rest/v1/{table}"
            f"?select={quote(select)}&order=id.asc&limit={page_size}&offset={offset}"
        )
        response = subprocess.check_output(
            [
                "curl",
                "-L",
                "-s",
                "--fail",
                url,
                "-H",
                f"apikey: {SERVICE_ROLE_KEY}",
                "-H",
                f"Authorization: Bearer {SERVICE_ROLE_KEY}",
            ],
            text=True,
        )
        page_rows = json.loads(response)
        rows.extend(page_rows)
        if len(page_rows) < page_size:
            break
        offset += page_size

    return rows


def as_text_list(values: list[str]) -> str:
    unique_values = []
    seen: set[str] = set()
    for value in values:
        normalized = (value or "").strip()
        if not normalized or normalized in seen:
            continue
        seen.add(normalized)
        unique_values.append(normalized)
    return "\n".join(unique_values)


def auto_fit_columns(sheet) -> None:
    for idx, column_cells in enumerate(sheet.columns, start=1):
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            longest_line = max((len(line) for line in value.splitlines()), default=0)
            max_length = max(max_length, min(longest_line, 120))
        sheet.column_dimensions[get_column_letter(idx)].width = min(max(max_length + 2, 12), 60)


def style_header(row) -> None:
    fill = PatternFill("solid", fgColor="163C33")
    for cell in row:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    regulations = fetch_json_pages("regulations", REGULATION_SELECT)
    source_documents = fetch_json_pages("regulation_source_documents", SOURCE_DOCUMENT_SELECT)

    docs_by_regulation: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for document in source_documents:
        docs_by_regulation[str(document.get("regulation_id") or "")].append(document)

    rows: list[dict[str, Any]] = []
    for regulation in regulations:
        regulation_id = str(regulation.get("id") or "")
        documents = docs_by_regulation.get(regulation_id, [])
        pdf_docs = [
            document
            for document in documents
            if (document.get("document_type") or "").lower() == "pdf"
            and (document.get("archived_public_url") or "").strip()
        ]
        if pdf_docs:
            continue

        rows.append(
            {
                "Regulation ID": regulation_id,
                "Title": regulation.get("title") or "",
                "Formal Title": regulation.get("formal_title") or "",
                "Region": regulation.get("region") or "",
                "Category": regulation.get("category") or "",
                "Status": regulation.get("status") or "",
                "Source Name": regulation.get("source_name") or "",
                "Source Link Kind": regulation.get("source_link_kind") or "",
                "Regulation Source URL": regulation.get("source_url") or "",
                "Official Website URL": regulation.get("official_source_url") or "",
                "Policy Page URL": regulation.get("policy_page_url") or "",
                "Regulation Updated At": regulation.get("updated_at") or "",
                "Source Document Count": len(documents),
                "PDF Source Document Count": len(
                    [document for document in documents if (document.get("document_type") or "").lower() == "pdf"]
                ),
                "Source Document Types": as_text_list(
                    [str(document.get("document_type") or "") for document in documents]
                ),
                "Source Document Version Labels": as_text_list(
                    [str(document.get("version_label") or "") for document in documents]
                ),
                "Document URLs": as_text_list([str(document.get("document_url") or "") for document in documents]),
                "Document Source URLs": as_text_list([str(document.get("source_url") or "") for document in documents]),
                "Document Official Website URLs": as_text_list(
                    [str(document.get("official_source_url") or "") for document in documents]
                ),
                "Document Policy Page URLs": as_text_list(
                    [str(document.get("policy_page_url") or "") for document in documents]
                ),
                "Archived PDF URLs": as_text_list(
                    [str(document.get("archived_public_url") or "") for document in documents]
                ),
            }
        )

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"

    summary_rows = [
        ("Generated At (UTC)", datetime.now(timezone.utc).isoformat()),
        ("Supabase Project URL", SUPABASE_URL),
        ("Total Regulations", len(regulations)),
        ("Total Source Documents", len(source_documents)),
        ("Regulations Without Archived Supabase PDF", len(rows)),
        (
            "Definition Used",
            "A regulation is in this file if it has no regulation_source_documents row with document_type=pdf and a non-empty archived_public_url.",
        ),
    ]

    for key, value in summary_rows:
        summary.append([key, value])

    for cell in summary["A"]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 36
    summary.column_dimensions["B"].width = 120
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet = workbook.create_sheet("No PDF Regulations")
    headers = list(rows[0].keys()) if rows else [
        "Regulation ID",
        "Title",
        "Formal Title",
        "Region",
        "Category",
        "Status",
        "Source Name",
        "Source Link Kind",
        "Regulation Source URL",
        "Official Website URL",
        "Policy Page URL",
        "Regulation Updated At",
        "Source Document Count",
        "PDF Source Document Count",
        "Source Document Types",
        "Source Document Version Labels",
        "Document URLs",
        "Document Source URLs",
        "Document Official Website URLs",
        "Document Policy Page URLs",
        "Archived PDF URLs",
    ]
    sheet.append(headers)
    style_header(sheet[1])

    for row in rows:
        sheet.append([row.get(header, "") for header in headers])

    sheet.freeze_panes = "A2"
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    auto_fit_columns(sheet)

    workbook.save(OUTPUT_PATH)
    print(json.dumps({"output": str(OUTPUT_PATH), "rows": len(rows)}, indent=2))


if __name__ == "__main__":
    main()
