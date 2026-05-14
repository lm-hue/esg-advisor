#!/usr/bin/env python3
from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


REPO_ROOT = Path("/Users/lucas/Documents/GitHub/esg-advisor")
TMP_ROOT = Path("/tmp")
OUTPUT_DIR = REPO_ROOT / "outputs"
OUTPUT_PATH = OUTPUT_DIR / f"regulation_button_links_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
OFFICIAL_ISSUER_FALLBACKS_PATH = REPO_ROOT / "src/lib/officialIssuerSiteFallbacks.ts"

ARCHIVED_SOURCE_PREFIX = (
    "https://twjaqynuamghrobhdasf.supabase.co/storage/v1/object/public/regulation-source-archives/"
)


def load_json_pages(prefix: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for index in range(3):
        path = TMP_ROOT / f"{prefix}_{index}.json"
        if path.exists():
            rows.extend(json.loads(path.read_text()))
    return rows


def parse_manual_fallbacks() -> dict[str, dict[str, str]]:
    text = (REPO_ROOT / "src/lib/manualSourceLinkFallbacks.ts").read_text()
    pattern = re.compile(r"'([0-9a-f-]{36})': \{ url: '([^']+)', kind: '([^']+)' \}")
    return {
        regulation_id: {"url": url, "kind": kind}
        for regulation_id, url, kind in pattern.findall(text)
    }


def parse_official_overrides() -> dict[str, str]:
    text = (REPO_ROOT / "src/lib/officialSourceUrlOverrides.ts").read_text()
    pattern = re.compile(r"'([0-9a-f-]{36})': '([^']+)'")
    return {regulation_id: url for regulation_id, url in pattern.findall(text)}


def parse_non_official_host_snippets() -> list[str]:
    text = OFFICIAL_ISSUER_FALLBACKS_PATH.read_text()
    match = re.search(r"const NON_OFFICIAL_HOST_SNIPPETS = \[(.*?)\]\n\nconst TRUSTED_OFFICIAL_HOSTS", text, re.S)
    if not match:
        return []
    return re.findall(r"'([^']+)'", match.group(1))


def parse_trusted_official_hosts() -> list[str]:
    text = OFFICIAL_ISSUER_FALLBACKS_PATH.read_text()
    match = re.search(r"const TRUSTED_OFFICIAL_HOSTS = \[(.*?)\]\n\nconst REGION_DEFAULT_URLS", text, re.S)
    if not match:
        return []
    return re.findall(r"'([^']+)'", match.group(1))


def parse_region_default_urls() -> dict[str, str]:
    text = OFFICIAL_ISSUER_FALLBACKS_PATH.read_text()
    match = re.search(r"const REGION_DEFAULT_URLS: Record<string, string> = \{(.*?)\n\}", text, re.S)
    if not match:
        return {}
    entries: dict[str, str] = {}
    for raw_key, bare_key, url in re.findall(r"\n\s+(?:'([^']+)'|([A-Za-z0-9_]+)): '([^']+)'", match.group(1)):
        key = raw_key or bare_key
        entries[key] = url
    return entries


def is_carrots_source_url(url: str | None) -> bool:
    return "carrotsandsticks.org" in (url or "").strip().lower()


def sanitize_frontend_source_url(url: str | None) -> str | None:
    normalized = (url or "").strip()
    if not normalized:
        return None
    if is_carrots_source_url(normalized):
        return None
    return normalized


def sanitize_frontend_source_name(name: str | None) -> str:
    normalized = (name or "").strip()
    if not normalized:
        return ""
    if re.search(r"carrots?\s*&\s*sticks|carrotsandsticks", normalized, re.IGNORECASE):
        return "Archived source"
    return normalized


def extract_host(url: str | None) -> str:
    normalized = (url or "").strip()
    if not normalized:
        return ""
    match = re.match(r"https?://([^/]+)", normalized, re.I)
    return match.group(1).lower() if match else ""


def is_trusted_official_publisher_url(
    url: str | None,
    non_official_host_snippets: list[str],
    trusted_official_hosts: list[str],
) -> bool:
    normalized = (url or "").strip().lower()
    if not normalized:
        return False
    host = extract_host(normalized)
    if not host:
        return False
    if any(snippet in host for snippet in non_official_host_snippets):
        return False
    if any(host == trusted or host.endswith(f".{trusted}") for trusted in trusted_official_hosts):
        return True
    return bool(
        host.endswith(".gov")
        or re.search(r"\.gov\.[a-z]{2}$", host)
        or "gouv." in host
        or "europa.eu" in host
        or "legislation." in host
        or "officialgazette" in host
        or "gazette" in host
        or "justice." in host
        or "statutebook" in host
        or "lovdata" in host
        or "e-tar" in host
        or "finlex" in host
        or "legilux" in host
        or "normattiva" in host
        or "gazzettaufficiale" in host
        or "boe.es" in host
        or "dre.pt" in host
        or "legifrance" in host
        or "bundesanzeiger" in host
        or "gesetze-im-internet" in host
        or "ifrs.org" in host
        or "globalreporting.org" in host
        or "oecd.org" in host
        or "sciencebasedtargets.org" in host
        or "tnfd.global" in host
        or "fsb-tcfd.org" in host
        or "cdp.net" in host
    )


def is_archived_source_url(url: str | None) -> bool:
    normalized = (url or "").strip().lower()
    return normalized.startswith(ARCHIVED_SOURCE_PREFIX.lower())


def looks_like_pdf_url(url: str | None) -> bool:
    normalized = (url or "").strip().lower()
    return "/api/download-pdf" in normalized or normalized.endswith(".pdf")


def resolve_source_link(
    record: dict[str, Any],
    official_overrides: dict[str, str],
    manual_fallbacks: dict[str, dict[str, str]],
    region_default_urls: dict[str, str],
    non_official_host_snippets: list[str],
    trusted_official_hosts: list[str],
) -> tuple[str | None, str | None]:
    fallback_source_link: tuple[str | None, str | None] | None = None

    current_official = sanitize_frontend_source_url(record.get("official_source_url"))
    if current_official:
        if is_trusted_official_publisher_url(current_official, non_official_host_snippets, trusted_official_hosts):
            return current_official, "official"
        kind = record.get("source_link_kind")
        if kind and kind != "official":
            fallback_source_link = (current_official, kind)
        elif is_archived_source_url(current_official):
            fallback_source_link = (current_official, "archived_pdf")
        else:
            fallback_source_link = (current_official, "reference_pdf" if looks_like_pdf_url(current_official) else "reference_page")

    override_url = official_overrides.get(record["id"])
    if override_url and is_trusted_official_publisher_url(override_url, non_official_host_snippets, trusted_official_hosts):
        return override_url, "official"

    source_url = sanitize_frontend_source_url(record.get("source_url"))
    if source_url:
        if is_trusted_official_publisher_url(source_url, non_official_host_snippets, trusted_official_hosts):
            return source_url, "official"
        if fallback_source_link is None:
            if is_archived_source_url(source_url):
                fallback_source_link = (source_url, "archived_pdf")
            else:
                fallback_source_link = (source_url, "reference_pdf" if looks_like_pdf_url(source_url) else "reference_page")

    manual = manual_fallbacks.get(record["id"])
    if manual and manual.get("kind") == "official":
        return manual["url"], manual["kind"]

    region_default = region_default_urls.get((record.get("region") or "").strip())
    if region_default:
        return region_default, "official"

    if manual:
        return manual["url"], manual["kind"]

    if fallback_source_link:
        return fallback_source_link

    return None, None


def normalize_regulation_record(
    record: dict[str, Any],
    official_overrides: dict[str, str],
    manual_fallbacks: dict[str, dict[str, str]],
    region_default_urls: dict[str, str],
    non_official_host_snippets: list[str],
    trusted_official_hosts: list[str],
) -> dict[str, Any]:
    source_url, source_kind = resolve_source_link(
        record,
        official_overrides,
        manual_fallbacks,
        region_default_urls,
        non_official_host_snippets,
        trusted_official_hosts,
    )
    return {
        **record,
        "source_name": sanitize_frontend_source_name(record.get("source_name")),
        "source_url": source_url or sanitize_frontend_source_url(record.get("source_url")) or "",
        "official_source_url": source_url,
        "source_link_kind": source_kind,
    }


def normalize_source_document_record(document: dict[str, Any]) -> dict[str, Any]:
    return {
        **document,
        "source_name": sanitize_frontend_source_name(document.get("source_name")),
        "source_url": sanitize_frontend_source_url(document.get("source_url")) or "",
        "official_source_url": sanitize_frontend_source_url(document.get("official_source_url")),
        "document_url": sanitize_frontend_source_url(document.get("document_url")),
        "archived_public_url": sanitize_frontend_source_url(document.get("archived_public_url")),
    }


def build_document_display_href(document: dict[str, Any]) -> str:
    if document.get("document_type") == "pdf":
        return (
            document.get("archived_public_url")
            or document.get("document_url")
            or document.get("official_source_url")
            or ""
        )
    return document.get("document_url") or document.get("official_source_url") or ""


def get_source_action_label(kind: str | None) -> str:
    return {
        "official": "View Official Source",
        "archived_pdf": "Open Archived PDF",
        "reference_pdf": "Open Source PDF",
        "reference_page": "Open Source Page",
    }.get(kind or "", "Open Source")


def autosize_columns(worksheet) -> None:
    for column in worksheet.columns:
        max_length = 0
        letter = get_column_letter(column[0].column)
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 70)


def style_header(worksheet, row_number: int = 1) -> None:
    fill = PatternFill("solid", fgColor="1F6E54")
    font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[row_number]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def main() -> None:
    regulations_raw = load_json_pages("regulations_live_page")
    source_documents_raw = load_json_pages("source_docs_live")

    official_overrides = parse_official_overrides()
    manual_fallbacks = parse_manual_fallbacks()
    region_default_urls = parse_region_default_urls()
    non_official_host_snippets = parse_non_official_host_snippets()
    trusted_official_hosts = parse_trusted_official_hosts()

    regulations = [
        normalize_regulation_record(
            record,
            official_overrides,
            manual_fallbacks,
            region_default_urls,
            non_official_host_snippets,
            trusted_official_hosts,
        )
        for record in regulations_raw
    ]
    source_documents = [
        normalize_source_document_record(document)
        for document in source_documents_raw
    ]

    documents_by_regulation: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for document in source_documents:
        documents_by_regulation[document["regulation_id"]].append(document)

    summary_rows: list[dict[str, Any]] = []
    source_button_rows: list[dict[str, Any]] = []
    pdf_button_rows: list[dict[str, Any]] = []

    for regulation in sorted(regulations, key=lambda row: (row.get("title") or "").lower()):
        regulation_id = regulation["id"]
        regulation_docs = documents_by_regulation.get(regulation_id, [])
        pdf_links: list[str] = []

        for document in regulation_docs:
            display_href = build_document_display_href(document)
            if not display_href:
                continue
            pdf_links.append(display_href)
            pdf_button_rows.append(
                {
                    "regulation_id": regulation_id,
                    "regulation_title": regulation.get("title", ""),
                    "detail_route": f"/framework-library/{regulation_id}",
                    "document_id": document.get("id", ""),
                    "document_type": document.get("document_type", ""),
                    "version_label": document.get("version_label", ""),
                    "source_name": document.get("source_name", ""),
                    "open_pdf_button_label": f"Open {'PDF' if document.get('document_type') == 'pdf' else 'Link'}",
                    "open_pdf_button_link": display_href,
                }
            )

        source_action_url = regulation.get("official_source_url")
        source_action_kind = regulation.get("source_link_kind")
        hide_primary_source_action = (
            source_action_kind == "archived_pdf"
            and bool(source_action_url)
            and any(build_document_display_href(document) == source_action_url for document in regulation_docs)
        )
        source_action_visible = bool(source_action_url) and not hide_primary_source_action
        source_action_label = get_source_action_label(source_action_kind) if source_action_visible else ""

        summary_rows.append(
            {
                "regulation_id": regulation_id,
                "title": regulation.get("title", ""),
                "detail_route": f"/framework-library/{regulation_id}",
                "source_name": regulation.get("source_name", ""),
                "open_source_button_visible": "Yes" if source_action_visible else "No",
                "open_source_button_label": source_action_label,
                "open_source_button_kind": source_action_kind or "",
                "open_source_button_link": source_action_url or "",
                "open_source_hidden_reason": "Hidden because it duplicates a versioned source-file PDF button"
                if hide_primary_source_action
                else "",
                "open_pdf_button_count": len(pdf_links),
                "open_pdf_button_links": "\n".join(pdf_links),
            }
        )

        if source_action_visible:
            source_button_rows.append(
                {
                    "regulation_id": regulation_id,
                    "regulation_title": regulation.get("title", ""),
                    "detail_route": f"/framework-library/{regulation_id}",
                    "source_name": regulation.get("source_name", ""),
                    "open_source_button_label": source_action_label,
                    "open_source_button_kind": source_action_kind or "",
                    "open_source_button_link": source_action_url or "",
                }
            )

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Regulation Summary"

    summary_sheet.append(
        [
            "Regulation ID",
            "Title",
            "Detail Route",
            "Source Name",
            "Open Source Visible",
            "Open Source Label",
            "Open Source Kind",
            "Open Source Link",
            "Open Source Hidden Reason",
            "Open PDF Button Count",
            "Open PDF Button Links",
        ]
    )
    for row in summary_rows:
        summary_sheet.append(
            [
                row["regulation_id"],
                row["title"],
                row["detail_route"],
                row["source_name"],
                row["open_source_button_visible"],
                row["open_source_button_label"],
                row["open_source_button_kind"],
                row["open_source_button_link"],
                row["open_source_hidden_reason"],
                row["open_pdf_button_count"],
                row["open_pdf_button_links"],
            ]
        )

    source_sheet = workbook.create_sheet("Open Source Buttons")
    source_sheet.append(
        [
            "Regulation ID",
            "Regulation Title",
            "Detail Route",
            "Source Name",
            "Button Label",
            "Button Kind",
            "Button Link",
        ]
    )
    for row in source_button_rows:
        source_sheet.append(
            [
                row["regulation_id"],
                row["regulation_title"],
                row["detail_route"],
                row["source_name"],
                row["open_source_button_label"],
                row["open_source_button_kind"],
                row["open_source_button_link"],
            ]
        )

    pdf_sheet = workbook.create_sheet("Open PDF Buttons")
    pdf_sheet.append(
        [
            "Regulation ID",
            "Regulation Title",
            "Detail Route",
            "Document ID",
            "Document Type",
            "Version Label",
            "Source Name",
            "Button Label",
            "Button Link",
        ]
    )
    for row in pdf_button_rows:
        pdf_sheet.append(
            [
                row["regulation_id"],
                row["regulation_title"],
                row["detail_route"],
                row["document_id"],
                row["document_type"],
                row["version_label"],
                row["source_name"],
                row["open_pdf_button_label"],
                row["open_pdf_button_link"],
            ]
        )

    overview_sheet = workbook.create_sheet("Overview", 0)
    overview_sheet.append(["Metric", "Value"])
    overview_sheet.append(["Generated At", datetime.now().isoformat(timespec="seconds")])
    overview_sheet.append(["Regulations Loaded", len(regulations)])
    overview_sheet.append(["Visible Open Source Buttons", len(source_button_rows)])
    overview_sheet.append(["Visible Open PDF Buttons", len(pdf_button_rows)])
    overview_sheet.append(
        [
            "Method",
            "Built from the current live-data snapshots in /tmp plus the exact current frontend button logic.",
        ]
    )

    for worksheet in workbook.worksheets:
        style_header(worksheet)
        autosize_columns(worksheet)
        worksheet.freeze_panes = "A2"
        for row in worksheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
