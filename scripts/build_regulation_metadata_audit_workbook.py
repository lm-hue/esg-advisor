from __future__ import annotations

import json
import re
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "outputs" / "regulation_metadata_audit_2026-04-26.xlsx"
LIVE_JSON_PATHS = [
    Path("/tmp/regulations_live_page_0.json"),
    Path("/tmp/regulations_live_page_1.json"),
    Path("/tmp/regulations_live_page_2.json"),
]
AUDIT_RESULTS_PATH = Path("/tmp/esg-link-audit-progress.json")
RESTORE_MANIFEST_PATH = Path("/tmp/carrots-pdf-restore-manifest.json")
GEOGRAPHY_TS_PATH = ROOT / "src" / "lib" / "geography.ts"
OFFICIAL_SOURCE_OVERRIDES_TS_PATH = ROOT / "src" / "lib" / "officialSourceUrlOverrides.ts"


JURISDICTION_TYPES = [
    {
        "value": "country",
        "label": "Country",
        "definition": "A national jurisdiction, such as Germany, Japan, or Canada.",
    },
    {
        "value": "supranational_region",
        "label": "Supranational region",
        "definition": "A rulemaking area above a single country, such as the European Union.",
    },
    {
        "value": "global",
        "label": "Global",
        "definition": "A globally applicable instrument or reference point that is not limited to one jurisdiction.",
    },
    {
        "value": "standards_body",
        "label": "Standards body",
        "definition": "An international body or initiative that publishes standards, frameworks, or guidance rather than country law.",
    },
    {
        "value": "exchange_or_regulator",
        "label": "Exchange / regulator",
        "definition": "A market operator or regulator that issues listing, disclosure, or supervisory rules.",
    },
]

STATUS_ORDER = [
    "proposal",
    "consultation",
    "draft",
    "adopted_not_yet_effective",
    "effective",
    "amended_effective",
    "repealed",
    "superseded",
    "archived",
]
STATUS_LABELS = {
    "proposal": "Proposal",
    "consultation": "Consultation",
    "draft": "Draft",
    "adopted_not_yet_effective": "Adopted",
    "effective": "In force",
    "amended_effective": "Amended in force",
    "repealed": "Repealed",
    "superseded": "Superseded",
    "archived": "Archived",
}
STATUS_DEFINITIONS = {
    "proposal": "An early-stage legislative or policy proposal that has not yet become formal draft text or adopted law.",
    "consultation": "A text or proposal open for public or stakeholder input before finalization.",
    "draft": "A draft measure, standard, or rule that is published but not yet adopted or effective.",
    "adopted_not_yet_effective": "A measure that has been formally adopted but whose operative date has not yet started.",
    "effective": "A measure currently in force and expected to apply now.",
    "amended_effective": "A measure currently in force that has been amended, with the amended version treated as the operative one.",
    "repealed": "A measure that has been formally revoked and is no longer operative.",
    "superseded": "A measure replaced in practice by a newer instrument or version.",
    "archived": "A historical or reference item kept for context, not for current operational reliance.",
}

REGULATION_TYPES = [
    ("law_or_regulation", "Law / regulation", "A binding law, regulation, directive, act, ordinance, or equivalent legal instrument."),
    ("proposal_or_draft", "Proposal / draft", "A proposed, draft, consultation, or exposure document that is not yet fully in force."),
    ("guidance", "Guidance", "Non-binding guidance, interpretive materials, manuals, FAQs, or implementation support."),
    ("standard", "Standard", "A formal standard with structured requirements, topics, metrics, or reporting criteria."),
    ("framework", "Framework", "A framework, protocol, or principles-based reference used to structure disclosure or ESG practice."),
    ("market_rule", "Market / exchange rule", "A rule or requirement issued by an exchange, securities authority, or market supervisor."),
    ("rating_or_benchmark", "Rating / benchmark", "A rating, benchmark, scoring, or ranking framework used to assess or compare entities."),
    ("policy_plan", "Policy / plan", "A roadmap, action plan, strategy, or programmatic policy document rather than a binding rule."),
]
REGULATION_TYPE_DEFINITIONS = {value: definition for value, _, definition in REGULATION_TYPES}
REGULATION_TYPE_LABELS = {value: label for value, label, _ in REGULATION_TYPES}

TOPICS = [
    ("reporting", "Reporting", "Disclosure, reporting, assurance, materiality, and sustainability statement requirements.", "report(ing), disclosure, assurance, materiality, ESRS, ISSB, IFRS, GRI"),
    ("taxonomy", "Taxonomy", "Classification systems that define what counts as sustainable or aligned activity.", "taxonomy"),
    ("governance", "Governance", "Board oversight, business conduct, ethics, anti-corruption, and governance controls.", "governance, board, anti-corruption, bribery, ethics, conduct"),
    ("human_rights", "Human rights", "Labour, workforce, equality, safety, harassment, and business and human rights topics.", "human rights, labour, worker, diversity, equality, harassment"),
    ("supply_chain", "Supply chain", "Due diligence, procurement, traceability, supplier oversight, and value-chain obligations.", "supply chain, due diligence, supplier, procurement, traceability"),
    ("biodiversity", "Biodiversity", "Nature, forests, ecosystems, land use, deforestation, and biodiversity impacts.", "biodiversity, deforestation, ecosystem, forest, nature-related"),
    ("water", "Water", "Water use, wastewater, marine impacts, and related aquatic resource issues.", "water, wastewater, marine, ocean"),
    ("pollution", "Pollution", "Air emissions, chemicals, contaminants, plastics, and pollution control obligations.", "pollution, emission, air quality, chemical, contaminant, plastic"),
    ("waste", "Waste", "Waste management, recycling, circularity, batteries, packaging, and end-of-life stewardship.", "waste, recycling, circular, battery stewardship, packaging"),
    ("energy", "Energy", "Energy use, efficiency, renewable energy, fuels, and energy transition topics.", "energy, electricity, renewable, efficiency, fuel"),
    ("climate", "Climate", "Climate risk, greenhouse gases, carbon, transition planning, and climate mitigation or adaptation.", "climate, greenhouse gas, GHG, net zero, carbon, TCFD"),
    ("finance", "Finance", "Sustainable finance, investor disclosure, funds, banks, securities, and financial-market rules.", "finance, financial, investor, bank, fund, securities, listing"),
]
TOPIC_DEFINITIONS = {value: definition for value, _, definition, _ in TOPICS}
TOPIC_LABELS = {value: label for value, label, _, _ in TOPICS}
TOPIC_KEYWORDS = {value: keywords for value, _, _, keywords in TOPICS}

DATE_PRECISION_OPTIONS = [
    ("year", "Year", "Only the year is reliable, so the displayed date should be treated as a year-level placeholder."),
    ("month", "Month", "The month and year are reliable, but a specific day is not confirmed."),
    ("day", "Day", "A full calendar date is available and can be shown precisely."),
]
DATE_PRECISION_DEFINITIONS = {value: definition for value, _, definition in DATE_PRECISION_OPTIONS}
DATE_PRECISION_LABELS = {value: label for value, label, _ in DATE_PRECISION_OPTIONS}

LINK_STATUS_DEFINITIONS = {
    "working_website": "The source URL resolves to a live website page rather than a downloadable PDF file.",
    "working_pdf": "The source URL resolves directly to a PDF document that appears usable as a source file.",
    "archived_pdf": "The original source was replaced with a locally archived PDF copy to preserve access.",
    "broken": "The recorded source URL is known to fail, redirect incorrectly, or no longer provide the expected source.",
    "missing": "No source URL is currently stored for the regulation.",
    "unknown": "The source URL exists but has not yet been confidently classified.",
}
LINK_STATUS_LABELS = {
    "working_website": "Working website",
    "working_pdf": "Working PDF",
    "archived_pdf": "Archived PDF",
    "broken": "Broken link",
    "missing": "No source link",
    "unknown": "Unknown",
}

SOURCE_HEALTH_DEFINITIONS = {
    "healthy": "The source link is currently usable without relying on an internal archive.",
    "archived": "The source is preserved through an archived PDF because the original link was unstable or unavailable.",
    "broken": "The source link is presently broken and should be replaced or restored.",
    "missing": "No source link is available yet.",
    "unknown": "Source health has not yet been confidently assessed.",
}
SOURCE_HEALTH_LABELS = {
    "healthy": "Healthy",
    "archived": "Archived",
    "broken": "Broken",
    "missing": "Missing",
    "unknown": "Unknown",
}

VOLUNTARY_FRAMEWORK_MATCHERS = [
    "ghg protocol",
    "wbcsd",
    "wri",
    "unep fi",
    "tnfd",
    "gri",
    "sasb",
    "sbti",
    "science based targets",
    "transition plan taskforce",
    "issb",
    "ifrs foundation",
    "ssbj",
    "integrated reporting",
    "equator principles",
    "un global compact",
    "principles for responsible investment",
    "principles for responsible banking",
    "principles for sustainable insurance",
]
STANDARD_MATCHERS = [
    "standard",
    "standards",
    "iso ",
    "isrs",
    "esrs",
    "ifrs s1",
    "ifrs s2",
    "sasb",
    "gri",
]
MARKET_RULE_MATCHERS = [
    "stock exchange",
    "listing requirement",
    "listing rule",
    "exchange rule",
    "hkex",
    "sgx",
    "jse",
    "sebi",
    "sec",
    "fca",
]
GUIDANCE_MATCHERS = ["guidance", "guide", "playbook", "faq", "manual"]
RATING_MATCHERS = ["rating", "ratings", "ranking", "rankings", "benchmark", "score", "cdp", "ecovadis", "msci"]
POLICY_PLAN_MATCHERS = ["roadmap", "policy plan", "action plan", "strategy", "programme", "program"]
PROPOSAL_MATCHERS = ["proposal", "proposed", "draft", "consultation", "exposure draft", "bill"]

THIN_BORDER = Border(
    left=Side(style="thin", color="D7DDD9"),
    right=Side(style="thin", color="D7DDD9"),
    top=Side(style="thin", color="D7DDD9"),
    bottom=Side(style="thin", color="D7DDD9"),
)
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F6F5D")
SUBHEADER_FILL = PatternFill(fill_type="solid", fgColor="EEF4F1")
ACCENT_FILL = PatternFill(fill_type="solid", fgColor="F7F2E3")


def yes_no(value: bool | None) -> str:
    if value is None:
        return "N/A"
    return "Yes" if value else "No"


def load_json(path: Path):
    return json.loads(path.read_text())


def load_live_rows(paths: Iterable[Path]) -> list[dict]:
    rows: list[dict] = []
    for path in paths:
        rows.extend(load_json(path))
    return rows


def load_geography_options(path: Path) -> dict[str, dict]:
    text = path.read_text()
    pattern = re.compile(
        r"\{\s*value:\s*'([^']+)'\s*,\s*name:\s*'([^']+)'\s*,\s*kind:\s*'([^']+)'\s*,\s*region:\s*'([^']+)'\s*,\s*continent:\s*'([^']+)'\s*,\s*emoji:\s*'([^']+)'",
        re.MULTILINE,
    )
    options: dict[str, dict] = {}
    for value, name, kind, region, continent, emoji in pattern.findall(text):
        options[value] = {
            "value": value,
            "name": name,
            "kind": kind,
            "region": region,
            "continent": continent,
            "emoji": emoji,
        }
    return options


def load_official_source_overrides(path: Path) -> dict[str, str]:
    text = path.read_text()
    pattern = re.compile(r"'([0-9a-f-]{36})': '([^']+)'")
    return {regulation_id: url for regulation_id, url in pattern.findall(text)}


def normalize_status(status: str | None) -> str:
    normalized = (status or "").strip().lower()
    if normalized == "proposal":
        return "proposal"
    if normalized == "consultation":
        return "consultation"
    if normalized == "draft":
        return "draft"
    if normalized in {"adopted_not_yet_effective", "adopted"}:
        return "adopted_not_yet_effective"
    if normalized in {"effective", "in_force"}:
        return "effective"
    if normalized in {"amended_effective", "amended"}:
        return "amended_effective"
    if normalized == "repealed":
        return "repealed"
    if normalized == "superseded":
        return "superseded"
    if normalized == "archived":
        return "archived"
    return "effective"


def normalize_category_key(category: str | None) -> str:
    normalized = (category or "").strip().lower()
    if normalized in {"environment", "environmental"}:
        return "Environmental"
    if normalized == "climate":
        return "Climate"
    if normalized == "circularity":
        return "Circularity"
    if normalized == "nature":
        return "Nature"
    if normalized == "social":
        return "Social"
    if normalized == "governance":
        return "Governance"
    return (category or "").strip()


def infer_regulation_type(row: dict) -> str:
    text = f"{row.get('title', '')} {row.get('source_name', '')}".lower()
    status = normalize_status(row.get("status"))
    if status in {"proposal", "consultation", "draft"} or any(matcher in text for matcher in PROPOSAL_MATCHERS):
        return "proposal_or_draft"
    if any(matcher in text for matcher in RATING_MATCHERS):
        return "rating_or_benchmark"
    if any(matcher in text for matcher in MARKET_RULE_MATCHERS):
        return "market_rule"
    if any(matcher in text for matcher in STANDARD_MATCHERS):
        return "standard"
    if any(matcher in text for matcher in GUIDANCE_MATCHERS):
        return "guidance"
    if any(matcher in text for matcher in POLICY_PLAN_MATCHERS):
        return "policy_plan"
    if any(matcher in text for matcher in VOLUNTARY_FRAMEWORK_MATCHERS):
        return "framework"
    return "law_or_regulation"


def infer_topics(row: dict) -> list[str]:
    text = " ".join(
        [
            row.get("title", "") or "",
            row.get("description", "") or "",
            row.get("full_description", "") or "",
            " ".join(row.get("tags", []) or []),
        ]
    ).lower()
    topics: list[str] = []

    def add(value: str) -> None:
        if value not in topics:
            topics.append(value)

    if re.search(r"\breport(ing)?\b|disclosure|disclosures|assurance|materiality|esrs|issb|ifrs|gri|sustainability report", text):
        add("reporting")
    if "taxonomy" in text:
        add("taxonomy")
    if re.search(r"governance|board|anti-corruption|bribery|ethics|conduct", text):
        add("governance")
    if re.search(r"human rights|labou?r|worker|forced labour|indigenous|harassment|diversity|equality", text):
        add("human_rights")
    if re.search(r"supply chain|due diligence|supplier|mineral|procurement|traceability", text):
        add("supply_chain")
    if re.search(r"biodiversity|deforestation|ecosystem|forest|nature-related", text):
        add("biodiversity")
    if re.search(r"\bwater\b|wastewater|marine|ocean", text):
        add("water")
    if re.search(r"pollution|emission|air quality|chemical|contaminant|plastic", text):
        add("pollution")
    if re.search(r"\bwaste\b|recycling|circular|battery stewardship|packaging", text):
        add("waste")
    if re.search(r"energy|electricity|renewable|efficiency|fuel", text):
        add("energy")
    if re.search(r"climate|greenhouse gas|ghg|net zero|carbon|tcfd|transition plan", text):
        add("climate")
    if re.search(r"finance|financial|investor|bank|fund|taxonomy|securities|listing", text):
        add("finance")

    if not topics:
        fallback = {
            "Climate": "climate",
            "Circularity": "waste",
            "Nature": "biodiversity",
            "Social": "human_rights",
            "Governance": "governance",
        }.get(normalize_category_key(row.get("category")))
        if fallback:
            add(fallback)

    if infer_regulation_type(row) == "standard":
        add("reporting")

    return topics


def infer_date_precision(row: dict) -> str | None:
    date_value = row.get("effective_date")
    if not date_value:
        return None
    parts = date_value.split("-")
    if len(parts) != 3:
        return "day"
    month, day = parts[1], parts[2]
    source_url = (row.get("source_url") or "").lower()
    if "carrotsandsticks.org" in source_url and month == "01" and day == "01":
        return "year"
    return "day"


def infer_jurisdiction_type(region: str | None, geography: dict[str, dict]) -> str:
    option = geography.get(region or "")
    if not option:
        return "global"
    if option["kind"] == "country":
        return "country"
    if option["kind"] == "region":
        return "supranational_region"
    if option["kind"] == "global":
        return "global"
    if option["value"] == "SSE":
        return "exchange_or_regulator"
    return "standards_body"


def implemented_link_status(source_url: str | None) -> str:
    url = (source_url or "").strip().lower()
    if not url:
        return "missing"
    if "/regulation-source-archives/" in url:
        return "archived_pdf"
    if url.endswith(".pdf"):
        return "working_pdf"
    if "carrotsandsticks.org" in url:
        return "broken"
    if url.startswith("http"):
        return "working_website"
    return "unknown"


def source_health_from_link_status(status: str | None) -> str:
    normalized = (status or "").strip().lower()
    if normalized in {"working_website", "working_pdf"}:
        return "healthy"
    if normalized == "archived_pdf":
        return "archived"
    if normalized == "broken":
        return "broken"
    if normalized == "missing":
        return "missing"
    return "unknown"


def audited_link_status(source_url: str | None, audit_map: dict[str, dict]) -> str:
    url = (source_url or "").strip()
    if not url:
        return "missing"
    audit = audit_map.get(url)
    if not audit:
        return "unknown"
    if audit.get("ok") and audit.get("kind") == "pdf":
        return "working_pdf"
    if audit.get("ok") and audit.get("kind") == "website":
        return "working_website"
    return "broken"


def extract_domain(url: str | None) -> str:
    if not url:
        return ""
    parsed = urlparse(url)
    return parsed.netloc.lower()


def looks_like_official_source_url(url: str | None) -> bool:
    normalized = (url or "").strip().lower()
    if not normalized:
        return False
    if "carrotsandsticks.org" in normalized:
        return False
    if "/regulation-source-archives/" in normalized:
        return False
    return normalized.startswith("http")


def frontend_surfaces_for_area(area: str) -> str:
    mapping = {
        "jurisdiction": "Regulation detail metadata card; existing region filters still use jurisdiction values.",
        "status": "Framework library filters, badges, timeline badges, detail page, glossary.",
        "dates": "Framework library cards/table/timeline, comparison views, detail page, glossary.",
        "regulation_type": "Framework library filters, badges, timeline badges, detail page, glossary.",
        "topics": "Framework library filters, badges, timeline badges, detail page, glossary.",
        "links": "Regulation detail link metadata card; workbook audit captures deeper link diagnostics.",
    }
    return mapping[area]


def set_title(ws, title: str, subtitle: str | None = None) -> int:
    ws["A1"] = title
    ws["A1"].font = Font(size=16, bold=True, color="1F2D27")
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(size=10, color="5B6B64")
        return 4
    return 3


def write_table(ws, start_row: int, headers: list[str], rows: list[list], freeze: bool = False) -> int:
    header_row = start_row
    for column_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, row in enumerate(rows, start=header_row + 1):
        for column_index, value in enumerate(row, start=1):
            cell = ws.cell(row=row_index, column=column_index, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}"
    if freeze:
        ws.freeze_panes = f"A{header_row + 1}"
    return header_row + len(rows) + 2


def autofit(ws, min_width: int = 12, max_width: int = 40) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        ws.column_dimensions[letter].width = max(min_width, min(max_width, max_length + 2))


def add_note_block(ws, row: int, heading: str, body: str) -> int:
    ws.cell(row=row, column=1, value=heading).font = Font(bold=True, color="1F2D27")
    ws.cell(row=row, column=1).fill = SUBHEADER_FILL
    ws.cell(row=row, column=1).border = THIN_BORDER
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    for column in range(1, 7):
        ws.cell(row=row, column=column).fill = SUBHEADER_FILL
        ws.cell(row=row, column=column).border = THIN_BORDER
    ws.cell(row=row + 1, column=1, value=body)
    ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=6)
    for column in range(1, 7):
        ws.cell(row=row + 1, column=column).border = THIN_BORDER
        ws.cell(row=row + 1, column=column).alignment = Alignment(wrap_text=True, vertical="top")
    return row + 3


def main() -> int:
    output_path = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else OUTPUT_PATH
    output_path.parent.mkdir(parents=True, exist_ok=True)

    live_rows = load_live_rows(LIVE_JSON_PATHS)
    audit_payload = load_json(AUDIT_RESULTS_PATH)
    audit_map = {result["url"]: result for result in audit_payload["results"]}
    restore_rows = load_json(RESTORE_MANIFEST_PATH)
    restore_by_regulation_id = {row["regulationId"]: row for row in restore_rows}
    geography = load_geography_options(GEOGRAPHY_TS_PATH)
    official_source_overrides = load_official_source_overrides(OFFICIAL_SOURCE_OVERRIDES_TS_PATH)

    normalized_status_counts = Counter()
    inferred_type_counts = Counter()
    inferred_topic_counts = Counter()
    inferred_precision_counts = Counter()
    jurisdiction_type_counts = Counter()
    implemented_link_counts = Counter()
    audited_link_counts = Counter()
    official_source_origin_counts = Counter()

    enriched_rows = []
    for row in live_rows:
        normalized_status = normalize_status(row.get("status"))
        inferred_type = infer_regulation_type(row)
        inferred_topics = infer_topics(row)
        inferred_precision = infer_date_precision(row)
        jurisdiction_type = infer_jurisdiction_type(row.get("region"), geography)
        app_link_status = implemented_link_status(row.get("source_url"))
        audit_link = audited_link_status(row.get("source_url"), audit_map)
        app_source_health = source_health_from_link_status(app_link_status)
        audit_source_health = source_health_from_link_status(audit_link)
        audit_result = audit_map.get(row.get("source_url", ""))
        restore_match = restore_by_regulation_id.get(row["id"])
        live_official_source_url = row.get("official_source_url")
        override_official_source_url = official_source_overrides.get(row["id"])
        if live_official_source_url:
            implemented_official_source_url = live_official_source_url
            official_source_origin = "live_db"
        elif override_official_source_url:
            implemented_official_source_url = override_official_source_url
            official_source_origin = "override_map"
        elif looks_like_official_source_url(row.get("source_url")):
            implemented_official_source_url = row.get("source_url")
            official_source_origin = "source_url"
        else:
            implemented_official_source_url = ""
            official_source_origin = "missing"

        normalized_status_counts[normalized_status] += 1
        inferred_type_counts[inferred_type] += 1
        jurisdiction_type_counts[jurisdiction_type] += 1
        if inferred_precision:
            inferred_precision_counts[inferred_precision] += 1
        for topic in inferred_topics:
            inferred_topic_counts[topic] += 1
        implemented_link_counts[app_link_status] += 1
        audited_link_counts[audit_link] += 1
        official_source_origin_counts[official_source_origin] += 1

        enriched_rows.append(
            {
                **row,
                "normalized_status": normalized_status,
                "inferred_type": inferred_type,
                "inferred_topics": inferred_topics,
                "inferred_precision": inferred_precision,
                "jurisdiction_type": jurisdiction_type,
                "jurisdiction_value": row.get("region") or "",
                "implemented_link_status": app_link_status,
                "implemented_source_health": app_source_health,
                "audited_link_status": audit_link,
                "audited_source_health": audit_source_health,
                "audit_http_status": audit_result.get("status") if audit_result else None,
                "audit_kind": audit_result.get("kind") if audit_result else "",
                "audit_content_type": audit_result.get("contentType") if audit_result else "",
                "audit_final_url": audit_result.get("finalUrl") if audit_result else "",
                "source_domain": extract_domain(row.get("source_url")),
                "implemented_official_source_url": implemented_official_source_url,
                "implemented_official_source_domain": extract_domain(implemented_official_source_url),
                "official_source_origin": official_source_origin,
                "recoverable_from_local_pdf": bool(restore_match),
                "local_pdf_path": restore_match.get("filePath") if restore_match else "",
            }
        )

    wb = Workbook()
    wb.remove(wb.active)

    summary_ws = wb.create_sheet("Executive Summary")
    next_row = set_title(
        summary_ws,
        "Regulation Metadata Audit",
        "Generated from local app code, live regulation exports, the saved link audit, and the saved Carrots & Sticks PDF restore manifest.",
    )
    rollout_rows = [
        [
            "jurisdiction_type / jurisdiction_value",
            "Yes",
            "Yes",
            "No",
            "Standardized in code and surfaced on the regulation detail page; filters still use jurisdiction values rather than a dedicated jurisdiction type filter.",
        ],
        [
            "Richer lifecycle statuses",
            "Yes",
            "Yes",
            "No",
            "Canonical lifecycle labels are normalized in code. Empty statuses are now hidden from the library and timeline filters.",
        ],
        [
            "published_date / adopted_date / date_precision",
            "Yes",
            "Yes",
            "No",
            "Detail page now shows published date, adopted date, and date precision. Library/timeline continue to use effective_date as the primary date for display.",
        ],
        [
            "First-class regulation_type",
            "Yes",
            "Yes",
            "No",
            "Used in filters, badges, glossary, and detail pages; currently inferred heuristically from title, source name, and status when not stored in data.",
        ],
        [
            "Controlled topics",
            "Yes",
            "Yes",
            "No",
            "Used in filters, badges, glossary, and detail pages; currently inferred heuristically from text and broad category fallbacks.",
        ],
        [
            "link_status / source_health",
            "Yes",
            "Yes",
            "No",
            "Visible on the detail page and included in this workbook. Current frontend values are inferred in code unless the database eventually stores audited fields.",
        ],
        [
            "official_source_url",
            "Yes",
            "Yes",
            "No",
            "Detail pages now prefer official_source_url for the main source CTA. Today it is supplied by code fallback and override mapping because the live database has not yet applied the new official_source_url migration.",
        ],
    ]
    next_row = write_table(
        summary_ws,
        next_row,
        ["Area", "Implemented in app code", "Visible in frontend", "Persisted in live DB", "Notes"],
        rollout_rows,
        freeze=True,
    )
    snapshot_rows = [
        ["Live regulations audited", len(live_rows), "Rows loaded from saved live exports."],
        ["Live regulation table has metadata columns", "No", "Saved API probe returned: column regulations.jurisdiction_type does not exist."],
        ["Selectable geography values in model", len(geography), "Parsed from src/lib/geography.ts."],
        ["Current normalized statuses present in live data", len([status for status, count in normalized_status_counts.items() if count > 0]), "Only populated statuses are now shown in lifecycle filter menus."],
        ["Unique source URLs in saved link audit", audit_payload["summary"]["totalUniqueUrls"], "From /tmp/esg-link-audit-progress.json."],
        ["Working unique URLs in saved link audit", audit_payload["summary"]["working"], "Unique URL count, not regulation row count."],
        ["Failing unique URLs in saved link audit", audit_payload["summary"]["failing"], "Unique URL count, not regulation row count."],
        ["Broken Carrots & Sticks rows recoverable from local PDFs", len(restore_rows), "Matched from /tmp/carrots-pdf-restore-manifest.json."],
        ["Rows with implemented official_source_url", sum(1 for row in enriched_rows if row["implemented_official_source_url"]), "Computed from live DB value, override map, or safe source_url fallback."],
        ["Rows with official_source_url from override map", official_source_origin_counts["override_map"], "Backfilled from Carrots source metadata and manual overrides in app code."],
    ]
    next_row = add_note_block(
        summary_ws,
        next_row,
        "Important rollout note",
        "The richer model is implemented in application code and frontend normalization, but the live Supabase regulations table still has the legacy schema. Until the migration is applied, the app infers status, regulation type, topics, jurisdiction type, date precision, link status, source health, and official_source_url on the fly.",
    )
    write_table(summary_ws, next_row, ["Metric", "Value", "Notes"], snapshot_rows)
    autofit(summary_ws, max_width=48)

    jurisdiction_ws = wb.create_sheet("Jurisdiction Types")
    next_row = set_title(
        jurisdiction_ws,
        "Jurisdiction Types",
        "These are the standardized jurisdiction types implemented in code. Jurisdiction values come from the shared geography model and remain visible in the frontend.",
    )
    jurisdiction_rows = [
        [
            item["value"],
            item["label"],
            item["definition"],
            jurisdiction_type_counts[item["value"]],
            "Yes",
            frontend_surfaces_for_area("jurisdiction"),
        ]
        for item in JURISDICTION_TYPES
    ]
    write_table(
        jurisdiction_ws,
        next_row,
        ["Value", "Label", "Definition", "Live regulation row count", "Visible in frontend", "Frontend surfaces"],
        jurisdiction_rows,
        freeze=True,
    )
    autofit(jurisdiction_ws, max_width=44)

    jurisdiction_values_ws = wb.create_sheet("Jurisdiction Values")
    next_row = set_title(
        jurisdiction_values_ws,
        "Jurisdiction Values",
        "Parsed from the shared geography model. These values standardize the selectable countries, regions, global bodies, and standards organizations used by the app.",
    )
    value_rows = []
    region_counts = Counter(row.get("region") or "" for row in live_rows)
    for option in sorted(geography.values(), key=lambda item: (item["continent"], item["name"])):
        jurisdiction_type = infer_jurisdiction_type(option["value"], geography)
        value_rows.append(
            [
                option["value"],
                option["name"],
                option["kind"],
                jurisdiction_type,
                option["region"],
                option["continent"],
                region_counts.get(option["value"], 0),
                "Yes",
            ]
        )
    write_table(
        jurisdiction_values_ws,
        next_row,
        ["Value", "Name", "Geography kind", "Mapped jurisdiction_type", "Parent region", "Continent", "Live regulation row count", "Used in frontend region selectors"],
        value_rows,
        freeze=True,
    )
    autofit(jurisdiction_values_ws, max_width=34)

    lifecycle_ws = wb.create_sheet("Lifecycle Statuses")
    next_row = set_title(
        lifecycle_ws,
        "Lifecycle Statuses",
        "Canonical lifecycle statuses are normalized in code. The filter UI now hides statuses that have zero current live records.",
    )
    lifecycle_rows = []
    empty_statuses = []
    for status in STATUS_ORDER:
        count = normalized_status_counts.get(status, 0)
        if count == 0:
            empty_statuses.append(STATUS_LABELS[status])
        lifecycle_rows.append(
            [
                status,
                STATUS_LABELS[status],
                STATUS_DEFINITIONS[status],
                count,
                "Yes" if count > 0 else "No",
                "Yes",
                "Yes",
            ]
        )
    next_row = write_table(
        lifecycle_ws,
        next_row,
        ["Value", "Label", "Definition", "Live normalized row count", "Shown in filters now", "Tooltip in filter and badges", "Included in glossary"],
        lifecycle_rows,
        freeze=True,
    )
    next_row = add_note_block(
        lifecycle_ws,
        next_row,
        "Empty lifecycle statuses removed from filters",
        f"Statuses with zero live records today: {', '.join(empty_statuses) if empty_statuses else 'None'}. They remain part of the canonical model, but they no longer clutter the filter menus until matching records exist.",
    )
    raw_status_rows = [[status, count] for status, count in Counter(row.get("status") for row in live_rows).most_common()]
    write_table(lifecycle_ws, next_row, ["Legacy live status value", "Current live row count"], raw_status_rows)
    autofit(lifecycle_ws, max_width=46)

    dates_ws = wb.create_sheet("Dates & Precision")
    next_row = set_title(
        dates_ws,
        "Dates and Precision",
        "The app uses effective_date as the primary timeline and library date. Published and adopted dates are now surfaced on the detail page, but the live database still lacks those columns until the migration is applied.",
    )
    field_rows = [
        ["effective_date", "Yes", "Yes", "Library cards, table, timeline, comparison views, detail page primary date card", "Yes", len([row for row in live_rows if row.get("effective_date")]), "Primary operative date used across the UI."],
        ["published_date", "Yes", "Yes", "Detail page metadata card", "No", "N/A", "Implemented in code and UI, but not yet persisted in the live regulations table."],
        ["adopted_date", "Yes", "Yes", "Detail page metadata card", "No", "N/A", "Implemented in code and UI, but not yet persisted in the live regulations table."],
        ["date_precision", "Yes", "Yes", "Library cards, timeline, comparison views, detail page metadata card, glossary", "No", "N/A", "Currently inferred in code for live legacy rows; year precision is used mainly for Carrots & Sticks placeholder dates."],
        ["official_source_url", "Yes", "Yes", "Regulation detail page main 'View Official Source' CTA", "No", "N/A", "Currently computed from a generated override map plus safe source_url fallback until the live DB migration is applied."],
    ]
    next_row = write_table(
        dates_ws,
        next_row,
        ["Field", "Implemented in code", "Visible in frontend", "Where visible", "Live DB column present", "Observed live non-empty count", "Notes"],
        field_rows,
        freeze=True,
    )
    usage_rows = [
        ["Framework library cards / table / timeline", "effective_date", "Falls back to updated_at or created_at for sorting/grouping if effective_date is missing."],
        ["Regulation detail page", "effective_date primary; published_date and adopted_date supplemental", "Date precision controls how dates render when the corresponding field exists."],
        ["Comparison views", "effective_date", "Displayed with the same date precision formatting helper."],
    ]
    next_row = add_note_block(
        dates_ws,
        next_row,
        "What date is used where",
        "The app still treats effective_date as the primary operative date for most user-facing chronology. Published and adopted dates are available for richer metadata context, especially on the detail page, but they are not yet the default date used for sorting or main timeline placement.",
    )
    next_row = write_table(dates_ws, next_row, ["Frontend surface", "Primary date used", "Behavior"], usage_rows)
    precision_rows = [
        [value, label, definition, inferred_precision_counts.get(value, 0)]
        for value, label, definition in DATE_PRECISION_OPTIONS
    ]
    write_table(dates_ws, next_row, ["Value", "Label", "Definition", "Inferred live row count"], precision_rows)
    autofit(dates_ws, max_width=52)

    type_ws = wb.create_sheet("Regulation Types")
    next_row = set_title(
        type_ws,
        "Regulation Types",
        "These types are first-class in the app model and frontend, but live rows currently inherit them through inference unless the migration and data backfill are applied.",
    )
    next_row = add_note_block(
        type_ws,
        next_row,
        "Official backing",
        "There is no single official external taxonomy encoded for regulation_type. The current implementation is an editorial normalization layer that uses title, source name, and lifecycle cues to classify instruments into consistent buckets for filtering and display.",
    )
    type_rows = [
        [
            value,
            REGULATION_TYPE_LABELS[value],
            REGULATION_TYPE_DEFINITIONS[value],
            inferred_type_counts.get(value, 0),
            "Yes",
            "Yes",
            "Heuristic classification from title, source name, and status.",
        ]
        for value, _, _ in REGULATION_TYPES
    ]
    write_table(
        type_ws,
        next_row,
        ["Value", "Label", "Definition", "Live inferred row count", "Visible in frontend", "Tooltip and glossary added", "Current logic"],
        type_rows,
        freeze=True,
    )
    autofit(type_ws, max_width=48)

    topics_ws = wb.create_sheet("Controlled Topics")
    next_row = set_title(
        topics_ws,
        "Controlled Topics",
        "Topics are controlled vocabulary values in the app model, but today they are inferred by keyword logic and broad-category fallbacks rather than a single official external taxonomy.",
    )
    next_row = add_note_block(
        topics_ws,
        next_row,
        "Clustering logic and official backing",
        "The current topic model is an internal editorial clustering system. It loosely aligns with common ESG disclosure domains seen in frameworks such as ESRS, IFRS sustainability standards, supply-chain due diligence instruments, and nature-related guidance, but the code does not claim a single official source of truth for the topic list. It uses keyword inference first and category fallback second.",
    )
    topic_rows = [
        [
            value,
            TOPIC_LABELS[value],
            TOPIC_DEFINITIONS[value],
            inferred_topic_counts.get(value, 0),
            TOPIC_KEYWORDS[value],
            "No single official taxonomy",
            "Yes",
            "Yes",
        ]
        for value, _, _, _ in TOPICS
    ]
    write_table(
        topics_ws,
        next_row,
        ["Value", "Label", "Definition", "Live inferred row count", "Keyword / clustering logic", "Official backing", "Visible in frontend", "Tooltip and glossary added"],
        topic_rows,
        freeze=True,
    )
    autofit(topics_ws, max_width=48)

    link_summary_ws = wb.create_sheet("Link Status Summary")
    next_row = set_title(
        link_summary_ws,
        "Link Status and Source Health",
        "This sheet separates the app's implemented link/source-health model from the saved HTTP audit so you can see the current frontend abstraction and the deeper source-level reality side by side.",
    )
    next_row = add_note_block(
        link_summary_ws,
        next_row,
        "Model vs audit",
        "The frontend currently derives link_status and source_health from simple URL-based rules unless the future database columns are populated. The saved HTTP audit is more reliable for actual link health. Both views are included below so the gap is explicit.",
    )
    implemented_summary_rows = [
        [LINK_STATUS_LABELS[key], implemented_link_counts.get(key, 0), LINK_STATUS_DEFINITIONS[key]]
        for key in ["working_website", "working_pdf", "archived_pdf", "broken", "missing", "unknown"]
    ]
    next_row = write_table(
        link_summary_ws,
        next_row,
        ["Implemented link_status", "Live regulation row count", "Definition"],
        implemented_summary_rows,
        freeze=True,
    )
    next_row = write_table(
        link_summary_ws,
        next_row,
        ["Audited link_status", "Live regulation row count", "Definition"],
        [[LINK_STATUS_LABELS[key], audited_link_counts.get(key, 0), LINK_STATUS_DEFINITIONS[key]] for key in ["working_website", "working_pdf", "archived_pdf", "broken", "missing", "unknown"]],
    )
    next_row = write_table(
        link_summary_ws,
        next_row,
        ["official_source_url origin", "Live regulation row count", "Meaning"],
        [
            ["Live DB", official_source_origin_counts["live_db"], "The official source URL was already present in the live database row export."],
            ["Override map", official_source_origin_counts["override_map"], "The official source URL comes from the generated Carrots source metadata override file."],
            ["Safe source_url fallback", official_source_origin_counts["source_url"], "The source_url already looked like a direct official/non-archived link, so the app treats it as the official source."],
            ["Missing", official_source_origin_counts["missing"], "No official source URL is currently known from live data or the Carrots source metadata."],
        ],
    )
    next_row = write_table(
        link_summary_ws,
        next_row,
        ["Source health", "Definition"],
        [[SOURCE_HEALTH_LABELS[key], SOURCE_HEALTH_DEFINITIONS[key]] for key in ["healthy", "archived", "broken", "missing", "unknown"]],
    )
    snapshot_rows = [
        ["Saved HTTP audit unique URLs completed", audit_payload["summary"]["completed"], ""],
        ["Saved HTTP audit unique URLs working", audit_payload["summary"]["working"], ""],
        ["Saved HTTP audit unique URLs failing", audit_payload["summary"]["failing"], ""],
        ["Recoverable broken Carrots & Sticks regulation rows from local PDFs", len(restore_rows), "These rows have a matched local PDF available but are not yet restored into the live dataset."],
    ]
    write_table(link_summary_ws, next_row, ["Metric", "Value", "Notes"], snapshot_rows)
    autofit(link_summary_ws, max_width=50)

    inventory_ws = wb.create_sheet("Link Inventory")
    next_row = set_title(
        inventory_ws,
        "Per-Regulation Link Inventory",
        "Each live regulation row with its implemented metadata view, audited link result, and local PDF recovery availability.",
    )
    headers = [
        "id",
        "title",
        "region",
        "raw_status",
        "normalized_status",
        "effective_date",
        "date_precision_inferred",
        "jurisdiction_type_inferred",
        "regulation_type_inferred",
        "topics_inferred",
        "source_name",
        "source_url",
        "source_domain",
        "implemented_official_source_url",
        "implemented_official_source_domain",
        "official_source_origin",
        "implemented_link_status",
        "implemented_source_health",
        "audited_link_status",
        "audited_source_health",
        "audit_http_status",
        "audit_kind",
        "audit_final_url",
        "recoverable_from_local_pdf",
        "local_pdf_path",
    ]
    rows = []
    for row in enriched_rows:
        rows.append(
            [
                row["id"],
                row.get("title"),
                row.get("region"),
                row.get("status"),
                STATUS_LABELS[row["normalized_status"]],
                row.get("effective_date"),
                DATE_PRECISION_LABELS.get(row.get("inferred_precision"), ""),
                next(item["label"] for item in JURISDICTION_TYPES if item["value"] == row["jurisdiction_type"]),
                REGULATION_TYPE_LABELS[row["inferred_type"]],
                ", ".join(TOPIC_LABELS[topic] for topic in row["inferred_topics"]),
                row.get("source_name"),
                row.get("source_url"),
                row.get("source_domain"),
                row.get("implemented_official_source_url"),
                row.get("implemented_official_source_domain"),
                row.get("official_source_origin"),
                LINK_STATUS_LABELS[row["implemented_link_status"]],
                SOURCE_HEALTH_LABELS[row["implemented_source_health"]],
                LINK_STATUS_LABELS[row["audited_link_status"]],
                SOURCE_HEALTH_LABELS[row["audited_source_health"]],
                row.get("audit_http_status"),
                row.get("audit_kind"),
                row.get("audit_final_url"),
                "Yes" if row["recoverable_from_local_pdf"] else "No",
                row.get("local_pdf_path"),
            ]
        )
    write_table(inventory_ws, next_row, headers, rows, freeze=True)
    for row_index in range(5, 5 + len(rows)):
        source_url_cell = inventory_ws.cell(row=row_index, column=12)
        if source_url_cell.value:
            source_url_cell.hyperlink = str(source_url_cell.value)
            source_url_cell.style = "Hyperlink"
        official_url_cell = inventory_ws.cell(row=row_index, column=14)
        if official_url_cell.value:
            official_url_cell.hyperlink = str(official_url_cell.value)
            official_url_cell.style = "Hyperlink"
        final_url_cell = inventory_ws.cell(row=row_index, column=23)
        if final_url_cell.value:
            final_url_cell.hyperlink = str(final_url_cell.value)
            final_url_cell.style = "Hyperlink"
    autofit(inventory_ws, max_width=42)

    meta_ws = wb.create_sheet("Workbook Notes")
    next_row = set_title(
        meta_ws,
        "Workbook Notes",
        "Small notes that matter when reading the audit.",
    )
    notes_rows = [
        ["Live schema status", "The saved Supabase API probe returned: column regulations.jurisdiction_type does not exist. That means the live database has not yet applied the richer metadata migration."],
        ["Frontend rollout status", "The frontend now exposes glossary entries and explanation hovers for lifecycle statuses, regulation types, and topics. The regulation detail page also exposes jurisdiction, published/adopted dates, date precision, link status, and source health."],
        ["Lifecycle filter cleanup", "Lifecycle filter menus now show only statuses that actually exist in the current loaded dataset, so empty statuses are removed from the user-facing filter lists."],
        ["Topic confidence", "Controlled topics are useful and standardized, but they are still heuristic. They should not be treated as an official external taxonomy without a future governance decision."],
        ["Link confidence", "The detailed inventory sheet includes both the implemented app-level link status and the saved HTTP audit result because those are not the same today."],
    ]
    write_table(meta_ws, next_row, ["Note", "Detail"], notes_rows, freeze=True)
    autofit(meta_ws, max_width=54)

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet["A1"].fill = ACCENT_FILL
        sheet["A1"].border = THIN_BORDER

    wb.save(output_path)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
